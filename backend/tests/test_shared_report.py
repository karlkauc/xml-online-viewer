"""Excel for a shared report, help slugs, and formula safety of the workbook."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.report.error_catalog import HELP, help_slug


def test_help_slugs_point_at_existing_pages() -> None:
    from app.content.pages import get_page

    for slug in HELP:
        assert get_page(f"errors/{slug}") is not None, slug
    assert help_slug("SCHEMAV_CVC_PATTERN_VALID") == "pattern-mismatch"
    missing = "Element 'R': Missing child element(s)."
    unexpected = "Element 'Z': This element is not expected."
    assert help_slug("SCHEMAV_ELEMENT_CONTENT", missing) == "missing-child-element"
    assert help_slug("SCHEMAV_ELEMENT_CONTENT", unexpected) == "element-not-expected"
    assert help_slug("SOMETHING_ELSE") is None


def test_shared_report_excel_keeps_formulas_as_text() -> None:
    client = TestClient(app)
    payload = {
        "xml_filename": "a.xml",
        "errors": [
            {
                "line": 3,
                "message": '=HYPERLINK("http://evil","x")',
                "severity": "error",
                "help_slug": "pattern-mismatch",
            }
        ],
    }
    r = client.post("/api/validate/report/excel", json=payload)
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    cells = [c for row in ws.iter_rows() for c in row if c.value == '=HYPERLINK("http://evil","x")']
    assert cells and all(c.data_type == "s" for c in cells)
    links = [c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink]
    assert links == ["https://www.xml-viewer.online/errors/pattern-mismatch"]


def test_shared_report_is_capped() -> None:
    client = TestClient(app)
    errors = [{"message": "x", "severity": "error"}] * 501
    assert client.post("/api/validate/report/excel", json={"errors": errors}).status_code == 422
